import re
import sys

from docx import Document


if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")


KEYWORDS = re.compile(
    r"인물|캐릭터|도단아|도수아|도영진|고혜정|하예진|이성호|고명준|김주호|"
    r"루시|도현서|이하랑|혜정|예진|성호|명준|주호|현서|하랑"
)


def heading_level(paragraph):
    style_name = paragraph.style.name if paragraph.style else ""
    match = re.match(r"Heading\s+(\d+)", style_name, re.IGNORECASE)
    return int(match.group(1)) if match else None


def extract(path):
    doc = Document(path)
    print(f"\n===== {path} =====")

    active_level = None
    for index, paragraph in enumerate(doc.paragraphs):
        text = paragraph.text.strip()
        if not text:
            continue
        level = heading_level(paragraph)
        if level is not None:
            if active_level is not None and level <= active_level:
                active_level = None
            if KEYWORDS.search(text):
                active_level = level
        if active_level is not None or KEYWORDS.search(text):
            print(f"P{index:04d} [{paragraph.style.name}] {text}")

    for table_index, table in enumerate(doc.tables):
        rows = []
        for row in table.rows:
            cells = [cell.text.strip().replace("\n", " / ") for cell in row.cells]
            row_text = " | ".join(cells)
            rows.append(row_text)
        if any(KEYWORDS.search(row) for row in rows):
            print(f"\n-- TABLE {table_index} --")
            for row in rows:
                print(row)


def print_table_range(path, start, end):
    doc = Document(path)
    print(f"\n===== TABLES {start}-{end}: {path} =====")
    for table_index in range(start, min(end + 1, len(doc.tables))):
        print(f"\n-- TABLE {table_index} --")
        for row in doc.tables[table_index].rows:
            cells = [cell.text.strip().replace("\n", " / ") for cell in row.cells]
            print(" | ".join(cells))


def print_paragraph_range(path, start, end):
    doc = Document(path)
    print(f"\n===== PARAGRAPHS {start}-{end}: {path} =====")
    for index in range(start, min(end + 1, len(doc.paragraphs))):
        paragraph = doc.paragraphs[index]
        text = paragraph.text.strip()
        if text:
            print(f"P{index:04d} [{paragraph.style.name}] {text}")


if __name__ == "__main__":
    if len(sys.argv) >= 5 and sys.argv[1] == "--tables":
        start_index = int(sys.argv[2])
        end_index = int(sys.argv[3])
        for input_path in sys.argv[4:]:
            print_table_range(input_path, start_index, end_index)
    elif len(sys.argv) >= 5 and sys.argv[1] == "--paras":
        start_index = int(sys.argv[2])
        end_index = int(sys.argv[3])
        for input_path in sys.argv[4:]:
            print_paragraph_range(input_path, start_index, end_index)
    else:
        for input_path in sys.argv[1:]:
            extract(input_path)
